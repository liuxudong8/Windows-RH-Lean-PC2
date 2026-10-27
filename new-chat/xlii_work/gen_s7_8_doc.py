# -*- coding: utf-8 -*-
"""攻防 XLII · S7.8 落档（OMML 公式版）：(L1)(L2) 长线攻坚第一轮——真实定义件与 Abel–SW 组装真证。
母本：gen_s7_7_doc.py 样式链；公式 Word 原生 OMML（LaTeX→MathML→OMML）。
事实来源：本轮 Lean 工程执行记录（C:\\proj2\\QRH\\S7_1_Main.lean + main_l2_full.txt，2604 jobs 全绿）。
惯例（2026-10-27 起）：文档公式一律 OMML 原生对象；禁 LaTeX 源码文本；代码标识符保持文本。
"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from latex2mathml.converter import convert as _l2m
import mathml2omml as _m2o

TARGET = r"C:\Users\shtcl\Doubao\chats\2026-10-16\document\d1\路2攻防XLII-S7.8-(L1)(L2)长线攻坚第一轮-真实定义件与Abel-SW组装真证.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)

def set_fonts(style, east, ascii_="Times New Roman", size=10.5, bold=False, color=None):
    style.font.name = ascii_
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme"):
        if rf.get(qn(attr)):
            del rf.attrib[qn(attr)]
    rf.set(qn("w:ascii"), ascii_); rf.set(qn("w:hAnsi"), ascii_); rf.set(qn("w:eastAsia"), east)

n = doc.styles["Normal"]
set_fonts(n, "宋体", size=10.5)
pf = n.paragraph_format
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.line_spacing = 1.5
pf.space_before = Pt(0); pf.space_after = Pt(0)
ind = n.element.get_or_add_pPr().get_or_add_ind()
ind.set(qn("w:firstLineChars"), "200"); ind.set(qn("w:firstLine"), "420")

tstyle = doc.styles["Title"]
set_fonts(tstyle, "黑体", size=16, bold=True)
tstyle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
tstyle.paragraph_format.space_after = Pt(6)

h1 = doc.styles["Heading 1"]
set_fonts(h1, "黑体", size=12, bold=True, color=RGBColor(0, 0, 0))
h1.paragraph_format.space_before = Pt(12); h1.paragraph_format.space_after = Pt(6)
h1.paragraph_format.keep_with_next = True

h2 = doc.styles["Heading 2"]
set_fonts(h2, "黑体", size=10.5, bold=True, color=RGBColor(0, 0, 0))
h2.paragraph_format.space_before = Pt(6); h2.paragraph_format.space_after = Pt(3)
h2.paragraph_format.keep_with_next = True

cap = doc.styles["Caption"]
set_fonts(cap, "黑体", size=9, bold=False)
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_before = Pt(4); cap.paragraph_format.space_after = Pt(4)
cap.paragraph_format.keep_with_next = True

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

_HAT_RE = re.compile(
    r"<m:limUpp><m:e>(.*?)</m:e><m:lim>.*?<m:t>\^</m:t>.*?</m:lim></m:limUpp>", re.S)

def _fix_hat(xml):
    return _HAT_RE.sub(
        lambda m: '<m:acc><m:accPr><m:chr m:val="\u0302"/></m:accPr><m:e>%s</m:e></m:acc>' % m.group(1),
        xml)

def _omml(latex):
    out = _m2o.convert(_l2m(latex))
    out = _fix_hat(out)
    if "xmlns:m" not in out[:120]:
        out = out.replace("<m:oMath", '<m:oMath xmlns:m="%s"' % M_NS, 1)
    return out

def add_math_el(p, latex):
    elm = parse_xml(_omml(latex))
    p._p.append(elm)

def math_para(latex, space_before=None, space_after=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    xml = ('<m:oMathPara xmlns:m="%s"><m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr>%s</m:oMathPara>'
           % (M_NS, _omml(latex)))
    p._p.append(parse_xml(xml))
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p

def _apply_run_font(r, bold=False, font=None, size=None):
    if bold:
        r.font.bold = True
    if font or size:
        r.font.name = "Times New Roman"
        if size:
            r.font.size = Pt(size)
        rpr = r._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman")
        if font:
            rf.set(qn("w:eastAsia"), font)

def para(text, style=None, align=None, bold=False, font=None, size=None, indent_chars=0,
         space_before=None, space_after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if indent_chars:
        ppr = p._p.get_or_add_pPr(); ind2 = ppr.get_or_add_ind()
        ind2.set(qn("w:firstLineChars"), str(indent_chars * 100))
    tokens = re.split(r"(\$[^$]*\$|\^\{[^}]*\}|_\{[^}]*\})", text)
    for tk in tokens:
        if not tk:
            continue
        if tk.startswith("$") and tk.endswith("$") and len(tk) > 1:
            add_math_el(p, tk[1:-1])
        elif tk.startswith("^{"):
            r = p.add_run(tk[2:-1]); r.font.superscript = True
        elif tk.startswith("_{"):
            r = p.add_run(tk[2:-1]); r.font.subscript = True
        else:
            r = p.add_run(tk)
        _apply_run_font(r, bold=bold, font=font, size=size)
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p

def h1t(text):
    return doc.add_heading(text, level=1)

def h2t(text):
    return doc.add_heading(text, level=2)

def caption(text):
    return doc.add_paragraph(text, style="Caption")

def set_cell_borders(cell, top=8, bottom=8):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = tcPr.find(qn("w:tcBorders"))
    if tb is None:
        tb = OxmlElement("w:tcBorders"); tcPr.append(tb)
    for tag, sz in (("top", top), ("bottom", bottom)):
        el = tb.find(qn("w:" + tag))
        if el is None:
            el = OxmlElement("w:" + tag); tb.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0")

def style_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT, east="宋体", size=9):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    ppr = p._p.get_or_add_pPr()
    ind3 = ppr.get_or_add_ind(); ind3.set(qn("w:firstLineChars"), "0"); ind3.set(qn("w:firstLine"), "0")
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(size); r.font.bold = bold
    rpr = r._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman")
    rf.set(qn("w:eastAsia"), east)

def add_table(headers, rows, widths=None):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for r_ in t.rows:
        rp = r_._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); rp.append(cs)
    for j, htxt in enumerate(headers):
        c = t.rows[0].cells[j]
        set_cell_borders(c, top=8, bottom=4)
        style_cell(c, htxt, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, east="黑体")
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.rows[i + 1].cells[j]
            bottom = 8 if i == len(rows) - 1 else 0
            set_cell_borders(c, top=0, bottom=bottom)
            al = WD_ALIGN_PARAGRAPH.CENTER if (j == 0 or isinstance(val, str) or len(str(val)) <= 12) else WD_ALIGN_PARAGRAPH.RIGHT
            style_cell(c, str(val), align=al)
    if widths:
        for j, w in enumerate(widths):
            for r_ in t.rows:
                r_.cells[j].width = Cm(w)
    return t

def add_page_number_footer():
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"
    r.append(t); fld.append(r)
    p._p.append(fld)

add_page_number_footer()

# ---------- 标题块 ----------
para("路2攻防XLII—S7.8：(L1)(L2) 长线攻坚第一轮", style="Title",
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("（真实定义件 + Abel–SW 组装真证 + trivialTerm 收敛件）",
     align=WD_ALIGN_PARAGRAPH.CENTER, font="黑体", size=14, bold=True, space_after=4)
para("——《谱对偶论证》系列（d1）：S7.1 形式化证明推进", align=WD_ALIGN_PARAGRAPH.CENTER,
     font="楷体", size=12, space_after=4)
para("2026 年 10 月 27 日", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

# ---------- 摘要 ----------
para("摘要：", font="黑体", size=10.5, bold=True, indent_chars=0, space_after=0)
para("S7.1 形式化证明长线 (L1)(L2) 的第一轮攻坚落位，全量 lake build QRH 2604 jobs 全绿。"
     "(i) (L2) 半形式化落位完成：siegel_walfisz_chi5 由裸 axiom 升级为「真实素数项 primeSum + Abel 求和 + "
     "SW 部分和 + 权重/衰减 axiom 件」的组装结构；sum_abel_partial（Icc 口径部分和版 Abel）与 prime_sum_bound"
     "（量子项组装）为真证，SW 本体仍为命名 axiom（PNT-in-AP 量级，下一步击破目标）。(ii) (L1) 真实定义件："
     "zeroWeightedSum / primePowerSum / trivialTerm 由 opaque 升级为真实定义（Finset 有限和 + Real.rpow / "
     "∑' 实级数 / 复和取实部），logDerivTerm 保持 opaque（导数件）。(iii) (L1b) 真证两件：(L1b-A) r=1/2 "
     "代数桥 $n^{-1/2}=(\\sqrt n)^{-1}$（衔接 (L2) primeSum K≡1 特例）；(L1b-B) trivialTerm 级数收敛件"
     "（x>1、r>0 可和：rpow 分解 + inv 界 + 几何级数比较）。(iv) 如实挂账：x=1 边界调和型发散、零点无条件"
     "收敛挂 N(T) 计数（(L1c) 长线）。", space_after=6)

para("关键词：Landau 显式公式；Abel 求和；Siegel–Walfisz；trivialTerm 收敛；Real.rpow；OMML",
     space_after=10)

# ---------- 1 概览 ----------
h1t("1　概览：两线并进与编译状态")
para("本轮把击破清单仅剩的两条长线 (L1)(L2) 各自推进第一轮：以「半形式化落位」口径（同 (L3)(L4) 模式）"
     "把空洞 axiom 升级为真实定义 + 组装真证 + 机制件显式化。编译证据：lake build QRH 2604 jobs 全绿"
     "（main_l2_full.txt，exit=0）。")
caption("表 1　本轮改动与编译状态")
hdr1 = ["模块", "本轮动作", "状态"]
rows1 = [
    ["Main（S7_1_Main.lean）", "(L2) 组装真证：sum_abel_partial + prime_sum_bound；(L1) 真实定义件；(L1b) 两件真证", "编译绿"],
    ["Main（import 层）", "新增 Pow.Real / SpecificLimits.Basic / InfiniteSum.ENNReal", "编译绿"],
    ["全量 lake build QRH", "2604 jobs，零错误", "全绿"],
]
add_table(hdr1, rows1, widths=[4.2, 8.2, 3.0])
para("注：编译咒语 = Git/elan PATH 前置 + lake build QRH；证据文件 C:\\proj2\\main_l2_full.txt。",
     size=9, indent_chars=0, space_after=6)

# ---------- 2 (L2) 半形式化落位 ----------
h1t("2　(L2) 半形式化落位：Abel–SW 组装真证")
h2t("2.1　primeSum：真实素数项")
para("素数项（核参数化）的真实定义——S7.1 §2 的素数幂加权项，K 为核、δ = logT/(2π)：")
math_para(r"\mathrm{primeSum}(K,X,\delta)=\sum_{n\le X}\Lambda(n)\chi_5(n)\cdot(\sqrt n)^{-1}\cdot K(\delta\log n)",
          space_after=4)
h2t("2.2　sum_abel_partial：Icc 口径部分和版 Abel（真证）")
para("部分和口径的 Abel 求和恒等式（A(n) 为 SW 部分和）：")
math_para(r"\sum_{n\le X} a_n b_n = A(X)\,b(X) - \sum_{n\le X-1} A(n)\,(b(n+1)-b(n)),\qquad A(n)=\sum_{i\le n}\Lambda(i)\chi_5(i)",
          space_after=4)
para("证明链：Icc 1 N 的分解恒等式（$N=\\lfloor X\\rfloor_+$）用 Finset.Subset.antisymm + Nat.le_pred_of_lt / "
     "Nat.sub_le 原子构造（omega 对截断减 N−1 有盲区，报 c−d≥2 类反例——显式截断减路径）；succ 归纳步走 "
     "insert 路线（新增元素在左，sum_insert），mul_sub 展开后 ring 收尾。v4.34 实证：Finset.sum_Icc_succ_top "
     "带 N≥1 前提，无前提 rw 必败——弃用。", )
h2t("2.3　prime_sum_bound：量子项组装（真证）")
para("组装结论（C = c·M·(C₂+C₃)，c 为 SW 常数、M 为权重上界、C₂/C₃ 为衰减件常数）：")
math_para(r"|\mathrm{primeSum}(K,X,\delta)|\le c\cdot M\cdot(C_2+C_3)",
          space_after=4)
para("证明链：rw [hps] 改 simpa using sum_abel_partial（避开 AddLeftMono stuck）；abs_add_le 用 "
     "(abs_add_le _ _).trans (by rw [abs_neg])；Finset.single_le_sum 元变量未锚定时先给显式 hf 再 simpa；"
     "根式恒等式 $\\sqrt N\\,(\\sqrt N)^{-1}=1$ 与 $n\\cdot(\\sqrt n)^{-3}=(\\sqrt n)^{-1}$ 用 field_simp [hne] + "
     "conv 定点 rw [← hsq]（Real.sq_sqrt）——全匹配 rw 会打穿 sqrt 内部产生高次怪物项，Lean 提示改用 ring_nf；"
     "rw [← hrootn] 同样换 conv_rhs（否则匹配到幂内部产生 $n^4(\\sqrt n)^{-9}$ 怪物项）。")
h2t("2.4　剩余挂账（设计边界）")
para("siegel_walfisz_chi5 仍为命名 axiom，但其内部结构已由三件显式化：prime_weight_bounds（单常数 M 兼 "
     "$|b_n|\\le Mn^{-1/2}$ 与 $|\\Delta b_n|\\le Mn^{-3/2}$）、sum_weight_decay（C₂）、sum_sqrt_decay（C₃）。"
     "SW 本体真证 = PNT-in-AP 量级长线——半形式化落位的终点即把裸 axiom 升级为「真实素数项 + Abel 组装」，"
     "下一步击破 SW 部分和本身。")

# ---------- 3 (L1) 真实定义件 ----------
h1t("3　(L1) 真实定义件：opaque → def")
para("L1（landau_explicit_L，Yildirim 公式 (6) 重排式）的组成项此前为 opaque 抽象算子——opaque 上无「真证」"
     "可言。第一步（同 (L2) 模式）：升级为真实定义。数学内容（Main.lean 注释，Yildirim 1991 公式 (6)，"
     "对 L(s,χ) 无条件）：")
math_para(r"\sum_{\rho}x^{\rho-r}/(\rho-r)=\sum_{\ell}x^{-(2\ell+r+a)}/(2\ell+r+a)-\sum'_{n\le x}\Lambda(n)\chi(n)n^{-r}-(L'/L)(r,\chi)",
          space_after=4)
para("χ₅ 为模 5 实特征（偶特征，局部常数 a=0，已代入 trivialTerm）。三个真实定义：")
h2t("3.1　primePowerSum：Finset 有限和 + Real.rpow")
math_para(r"\mathrm{primePowerSum}(\chi,r,x)=\sum_{n\in[1,\lfloor x\rfloor]}\Lambda(n)\chi(n)\cdot n^{-r}",
          space_after=4)
h2t("3.2　trivialTerm：∑' 实级数（a=0）")
math_para(r"\mathrm{trivialTerm}(r,x)=\sum_{\ell=0}^{\infty}x^{-(2\ell+r)}/(2\ell+r)",
          space_after=4)
h2t("3.3　zeroWeightedSum：复和取实部")
math_para(r"\mathrm{zeroWeightedSum}(Z,r,x)=\mathrm{Re}\sum_{\rho\in Z}x^{\rho-r}/(\rho-r)",
          space_after=4)
para("零点和的无条件收敛需 N(T) 计数估计——挂 (L1c) 长线；前提 r ∉ Z 保证无 ρ=r 项（除零口径由 Lean "
     "0/0=0 兜底）。logDerivTerm 保持 opaque（导数件：deriv LFun / LFun，下一轮）。")

# ---------- 4 (L1b) 真证两件 ----------
h1t("4　(L1b) 真证两件")
h2t("4.1　(L1b-A) r=1/2 代数桥")
para("素数项桥恒等式（n ≥ 1）：")
math_para(r"n^{-1/2}=(\sqrt n)^{-1}",
          space_after=4)
para("证明链：Real.rpow_neg（$x^{-y}=(x^y)^{-1}$，0≤x）+ Real.sqrt_eq_rpow（$\\sqrt x=x^{1/2}$）。由此：")
math_para(r"\mathrm{primePowerSum}(\chi,1/2,X)=\sum_{n\le X}\Lambda(n)\chi(n)\cdot(\sqrt n)^{-1}",
          space_after=4)
para("——即 (L1) 的 r=1/2 特例与 (L2) primeSum（K≡1 时）精确衔接。v4.34 实证：rw 对 Real.rpow 的 ^ 记法"
     "匹配不可靠（`↑n ^ (-?y)` 匹配失败）→ 显式传参 exact/calc。")
h2t("4.2　(L1b-B) trivialTerm 收敛件")
para("定理：x > 1、r > 0 时 trivialTerm 级数可和。证明三步：")
para("(i) rpow 分解（x>0）：")
math_para(r"x^{-(2\ell+r)}=x^{-r}\cdot(x^{-2})^{\ell}",
          space_after=4)
para("——rpow_add（$x^{y+z}=x^y x^z$）+ rpow_mul（$x^{(-2)\ell}=(x^{-2})^{\ell}$）+ rpow_natCast 转 Nat 幂。")
para("(ii) 分母界：$2\\ell+r\\ge r>0$，故 $(2\\ell+r)^{-1}\\le r^{-1}$（inv_le_inv₀）。")
para("(iii) 逐项上界 + 几何级数：$f_\\ell\\le x^{-r}(x^{-2})^{\\ell}/r$，而 $x>1$ 时 $0\\le x^{-2}<1$"
     "（x^{-2} = (x^{-1})^2，inv 乘法链证 x^{-1}<1），故 $\\sum_\\ell (x^{-2})^{\\ell}$ 收敛"
     "（summable_geometric_of_lt_one）；Summable.of_nonneg_of_le 比较判别收尾。")
para("边界如实记录：x = 1 时 x^{−(2ℓ+r)} = 1，级数呈调和型发散——landau_explicit_L 的 1 ≤ x 前提在 x=1 "
     "处需单独口径（axiom 注释已注明）。")

# ---------- 5 挂账与边界 ----------
h1t("5　挂账与边界（如实记录）")
caption("表 2　(L1)(L2) 攻坚状态快照（2026-10-27）")
hdr2 = ["项", "状态"]
rows2 = [
    ["(L2) sum_abel_partial（Abel 部分和）", "真证（本档 §2.2）"],
    ["(L2) prime_sum_bound（量子项组装）", "真证（本档 §2.3）"],
    ["(L2) siegel_walfisz_chi5（SW 本体）", "命名 axiom；内部三件显式化——PNT-in-AP 量级长线"],
    ["(L1) 真实定义件（primePowerSum/trivialTerm/zeroWeightedSum）", "opaque → def（本档 §3）"],
    ["(L1b-A) r=1/2 代数桥", "真证（本档 §4.1）"],
    ["(L1b-B) trivialTerm 收敛（x>1、r>0）", "真证（本档 §4.2）；x=1 边界调和发散已注明"],
    ["(L1) logDerivTerm", "opaque——导数件，下一轮"],
    ["(L1) 零点无条件收敛", "挂 (L1c)——N(T) 计数估计长线"],
]
add_table(hdr2, rows2, widths=[10.0, 5.4])
para("下一步候选：(a) logDerivTerm 真实定义与导数件（deriv LFun）；(b) siegel_walfisz_chi5 本体击破"
     "（PNT-in-AP，挂 Ben Green 5 页证明的机制核验）；(c) 本轮变更推 GitHub（先更新 QRH_summary.md / "
     "QRH_summary_en.md，待用户指示）。", space_after=6)

# ---------- 参考文献 ----------
h1t("参考文献")
refs = [
    "[1] 攻防 XLII-S7.7（d1 已交付）：击破清单四连——BGS 引理 3–5 移植与 N1/L1/L4 真证（2026-10-27）.",
    "[2] 攻防 XLII-S7.6（d1 已交付）：SW 常数显式化——L(s,χ₅) 素数项显式界 (C,c).",
    "[3] 攻防 XLII-S7.1（d1 已交付）：形式化证明 Lean 落地——Yildirim 公式 L 版展开与 BGS 引理逐条移植（2026-10-25）.",
    "[4] 本工作 Lean 工程（2026-10-27）：C:\\proj2\\QRH\\S7_1_Main.lean（(L1)(L2) 区块，v4.34.0-rc2 + mathlib）；"
        "编译证据 main_l2_full.txt（2604 jobs 全绿）.",
    "[5] 本工作文档（2026-10-27）：QRH_summary.md / QRH_summary_en.md（精确陈述口径）.",
]
for rtxt in refs:
    para(rtxt, size=9, indent_chars=0, space_after=2)

doc.save(TARGET)
print("saved:", TARGET)

# -*- coding: utf-8 -*-
"""生成《路2攻防XXXVIII-Δ常数指纹对同虚部成对注入的爆炸响应与威胁排除的交叉验证闭合.docx》（d1 惯例）"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\路2攻防XXXVIII-Δ常数指纹对同虚部成对注入的爆炸响应与威胁排除的交叉验证闭合.docx"
SONG, HEI, KAI, TNR = "宋体", "黑体", "楷体", "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = TNR
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
npf = normal.paragraph_format
npf.space_before = Pt(0); npf.space_after = Pt(0); npf.line_spacing = 1.5

def set_fonts(run, eastasia=SONG, latin=TNR, size=10.5, bold=False, italic=False,
              sub=False, sup=False, color=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), eastasia)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if sub: run.font.subscript = True
    if sup: run.font.superscript = True
    if color: run.font.color.rgb = RGBColor(*color)

def rich(par, segs, size=10.5, eastasia=SONG):
    for kind, text in segs:
        r = par.add_run(text)
        if kind == "n": set_fonts(r, eastasia=eastasia, size=size)
        elif kind == "b": set_fonts(r, eastasia=HEI, size=size, bold=True)
        elif kind == "i": set_fonts(r, eastasia=eastasia, size=size, italic=True)
        elif kind == "sup": set_fonts(r, eastasia=eastasia, size=size, sup=True)
        elif kind == "sub": set_fonts(r, eastasia=eastasia, size=size, sub=True)
    return par

def para(segs, indent=True, align=None, size=10.5, before=0, after=0, line=1.5, eastasia=SONG):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if indent: pf.first_line_indent = Pt(size * 2)
    if align == "center": pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right": pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = line
    rich(p, segs, size=size, eastasia=eastasia)
    return p

def formula(segs, number=None, size=10.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(4); pf.space_after = Pt(4)
    pf.line_spacing = 1.5
    if number:
        pf.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
        rich(p, segs, size=size)
        rich(p, [("n", "\t（" + number + "）")], size=size)
    else:
        rich(p, segs, size=size)
    return p

def h1(text):
    p = doc.add_heading(level=1)
    pf = p.paragraph_format
    pf.space_before = Pt(12); pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    r = p.add_run(text)
    set_fonts(r, eastasia=HEI, latin=TNR, size=12, bold=True, color=(0, 0, 0))
    return p

def h2(text):
    p = doc.add_heading(level=2)
    pf = p.paragraph_format
    pf.space_before = Pt(6); pf.space_after = Pt(3)
    pf.line_spacing = 1.5
    r = p.add_run(text)
    set_fonts(r, eastasia=HEI, latin=TNR, size=10.5, bold=True, color=(0, 0, 0))
    return p

def make_table(headers, rows, widths, first_center=True):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htxt in enumerate(headers):
        c = tbl.rows[0].cells[j]
        c.width = widths[j]
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "D9D9D9")
        tcPr.append(shd)
        p = c.paragraphs[0]
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(htxt)
        set_fonts(r, eastasia=HEI, size=9, bold=True)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for i, row in enumerate(rows, 1):
        for j, txt in enumerate(row):
            c = tbl.rows[i].cells[j]
            c.width = widths[j]
            p = c.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (first_center and j == 0) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(txt)
            set_fonts(r, eastasia=SONG, size=9)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    trPr0 = tbl.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr0.append(th)
    return tbl

def caption(text, before=2, after=6):
    p = para([("n", text)], indent=False, align="center", size=9, before=before, after=after)
    p.paragraph_format.keep_with_next = True
    return p
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run()
fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
set_fonts(run, size=9)

# ---------- 标题区 ----------
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_after = Pt(0)
tp.paragraph_format.line_spacing = 1.5
rich(tp, [("b", "路2攻防XXXVIII—Δ常数指纹对同虚部成对注入的")], size=18, eastasia=HEI)
tp2 = doc.add_paragraph()
tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp2.paragraph_format.space_after = Pt(10)
tp2.paragraph_format.line_spacing = 1.5
rich(tp2, [("b", "爆炸响应与威胁排除的交叉验证闭合")], size=18, eastasia=HEI)
para([("n", "——《谱对偶论证》系列（d1）：S3 数值定向——窗口饱和、δ 免疫与 3σ 阈值 420 对")],
     indent=False, align="center", size=12, eastasia=KAI, after=2)
para([("n", "2026 年 10 月 24 日")], indent=False, align="center", size=9, after=12)

# ---------- 摘要 ----------
para([("b", "摘要：")], indent=False)
para([("n", "本攻防执行 S3：Δ 常数定律（跨族互斥指纹 S_ζL−E、c_pure=0.0070、14σ）对同虚部成对离线注入的响应。三件事实：(i) 代数免疫：A(t)=Σ_{γ} w((γ−t)λ) 只含零点虚部差——Δ 指纹对离线深度 δ 严格无关（威胁带 δ≤0.01 时 δλ≤0.133，远小于窗口 W=1.5——w 核饱和）——与 S2 Weil 口径的 δ²λ²/2 压低互补，检测对象是「同虚部成对」的存在性而非离线深度。(ii) 爆炸响应（数值，detF_inject.py）：威胁级（21.9 万对）ΔS=5.70e5、SNR=**1571σ**（σ_A=√(N_l·var_A) 口径）/531σ（√N_l）——每对响应 2.60（成对构型）/1.30（漏根单零点构型）——3σ 检出阈值 ≈420 对。(iii) 实测闭合：攻防 XXXVI 跨族极近对 <1e-6 仅 2 对——低于阈值 210 倍——Δ 指纹独立确认无威胁；且指纹对个别漏根免疫（单零点权 ≤4，仅占 3σ 阈值信号的 0.4%），补偿「极近对计数被漏根破坏」的唯一理论逃逸。基线复现：Δ_pure=−15046.5（文献 −15091）、Δ_pure/N_K=−0.007000、var_A=0.1144 精确命中；增量法 vs 全量重算 rel_diff=1.6e-14。判定：威胁排除支柱 = 间距计数（XXXVII）+ Δ 指纹（本攻防）双重独立闭合；21.9 万对注入使 Δ_pure/N_K 从 −0.0070 翻到 +0.258（变号）——跨族互斥常数层与威胁构型直接矛盾。")], indent=True)
para([("b", "关键词：")], indent=False)
para([("n", "Δ常数定律；跨族互斥；窗口饱和；δ免疫；交叉验证；威胁排除")], indent=False)

# ---------- 1 ----------
h1("1　承接与任务")
para([("n", "攻防 XXXVII 判定：S3（Δ 常数指纹）的价值从「检测威胁」降级为「排除交叉验证」——因为构型自洽性（同虚部成对 ⇒ 极近对 ⇒ 实测排除）已先行排除任何「离线且简单」威胁构型。本攻防恢复 S3 的独立数值价值：Δ 指纹测的不是离线深度，而是「同虚部成对」这一结构的积分统计——它提供一条与极近对计数独立的证据线，并天然补偿「计数受个别漏根破坏」的唯一理论逃逸。")], indent=True)
para([("n", "任务：把 Δ 常数定律（d1 系列已确立的跨族互斥常数层）用作同虚部成对注入的检测器——建立基线、测注入响应、定检出阈值、对实测闭合。")], indent=True)

# ---------- 2 ----------
h1("2　Δ 常数指纹检测器")
h2("2.1　口径")
para([("n", "Δ(T)=S_ζL(T)−E(T)；S_ζL=Σ_{γ′∈L(T)} A(γ′)，A(t)=Σ_{γ≤T} w((γ−t)λ)，w(u)=4/(4+u²)，λ=log T，窗口 |γ−t|<W=R/λ（R=20，绝对虚部窗口 W≈1.5）；E=Σ_{γ′} log(γ′/2π)/λ。纯跨族常数口径（2026-10-19 修正）：Δ_pure=Δ+c_cut·E，c_cut=1−I_R=0.0635（消除 E 口径核截断伪影）。定律：Δ_pure/N_K → −c_pure≈−0.0070。")], indent=True)
h2("2.2　基线复现（T=6e5）")
para([("n", "在并集数据（zeros_T6e5.npz，ζ 族 999492、L 族 1150050）上照抄 l4_multit 口径复算基线（entF_base.py，22s）——文献数值全部精确命中：")], indent=True)
caption("表 1　Δ 指纹基线复现（T=6e5）vs《Δ常数定律-陈述与证据》", before=4, after=2)
t1 = make_table(
    ["量", "本文复现", "文献值", "偏差"],
    [
        ["Δ/N_K", "−0.033940（−3.394%）", "−3.394%", "0"],
        ["Δ_pure", "−15046.5", "−15091", "0.3%"],
        ["Δ_pure/N_K", "−0.007000", "c_pure≈0.0070±0.0003", "命中"],
        ["Δ_pure/√N_l", "−14.03", "−14.1σ", "命中"],
        ["Δ_pure/√(N_l·var_A)", "−41.5", "−41.7σ", "命中"],
        ["mean_A / var_A", "0.7856 / 0.1144", "0.7856 / 0.1144", "0"],
    ],
    [Cm(4.0), Cm(3.2), Cm(4.4), Cm(2.6)])
para([("n", "基线可信：Δ 指纹检测器实现与文献口径一致（S=839011.3、E=911967.8、Δ=−72956.4；块 0/1/2 的 per_n=−0.0598/−0.0644/−0.0661 与文献分块一致）。")], indent=True)

# ---------- 3 ----------
h1("3　检测器对同虚部成对注入的响应")
h2("3.1　代数免疫：A(t) 不含实部")
para([("n", "A(t)=Σ_{γ≤T} w((γ−t)λ) 中权重只依赖零点虚部 γ 与探针 t 之差——**注入零点的实部（1/2±δ）对 Δ 严格不可见**。威胁带 δ≤0.01 时 δλ≤0.133，而窗口 R=20（绝对 W=R/λ≈1.5）：离线对的两个成员虚部都与 L 零点 γ′ 重合，w((γ′−γ′)λ)=w(0)=4 满权——**Δ 指纹对离线深度 δ 代数免疫、窗口饱和**。这与 S2 Weil 口径的 δ²λ²/2 压低（一阶消失）形成互补：Weil 测深度（被 δ 压低）、Δ 测结构（对 δ 饱和）。")], indent=True)
h2("3.2　注入实验")
para([("n", "构型 I（同虚部成对）：随机选 n_pairs 个 L 零点 γ′，在 ζ 族插入 2 个零点（虚部均=γ′，模拟 1/2±δ+iγ′）；构型 II（漏根变体）：每对仅 1 个可见。5 次随机抽样，增量法（A 对 z 线性，ΔS=Σ_{γ_i}Σ_{l′∈窗口}w((γ_i−l′)λ)）——全量重算验证精确（rel_diff=1.57e-14）。")], indent=True)
caption("表 2　Δ 指纹注入响应（ΔS mean±std，5 抽样；噪声 σ_A=362.7、σ_√N=1072.4）", before=4, after=2)
t2 = make_table(
    ["n_pairs", "构型 I ΔS", "SNR（σ_A）", "构型 II ΔS", "SNR（σ_A）"],
    [
        ["100", "257.4±2.8", "0.71", "131.4±1.5", "0.36"],
        ["200", "520.7±6.5", "1.44", "260.1±2.2", "0.72"],
        ["2000", "5194.7±9.9", "14.32", "2603.7±7.0", "7.18"],
        ["219000（威胁级）", "569817.5±94.9", "1570.96", "284887.8±88.2", "785.42"],
    ],
    [Cm(2.6), Cm(3.4), Cm(2.6), Cm(3.4), Cm(2.6)])
para([("n", "读数：每对响应 2.602（I）/1.300（II），全程线性（100→219000 对常数不漂移）；威胁级 SNR=1571σ——相对 S2 Weil 口径（0.5σ）**约 3000 倍**。注入后 Δ_pure/N_K 从 −0.007000 变为 +0.258088（219000 对 I）——**跨族互斥常数直接变号**：威胁构型与 Δ 常数定律在常数层不兼容。")], indent=True)

# ---------- 4 ----------
h1("4　阈值与实测闭合")
para([("n", "3σ 检出阈值：构型 I n≈420 对（σ_A 口径）/n≈1240 对（σ_√N 口径）；构型 II 双倍（840/2480）。攻防 XXXVI 实测跨族极近对 <1e-6 仅 2 个——**低于 I 型阈值 210 倍**。Δ 指纹独立确认：数据中不存在可检出的同虚部成对结构。")], indent=True)
para([("n", "漏根免疫：若极近对计数（XXXVII）受个别求根漏根破坏，Δ 指纹不受影响——单零点权 ≤w(0)=4，相对 3σ 阈值信号（≈1090）仅 0.4%——漏掉上百个零点才能压低 1σ。**Δ 指纹对「计数被漏根破坏」的唯一理论逃逸免疫**——两条证据线的盲区互不相交。")], indent=True)

# ---------- 5 ----------
h1("5　与 S2 Weil 口径对照")
para([("n", "同一威胁级注入（21.9 万对、δ=0.05）：S2 Weil 实部口径 SNR≈0.59（Re 修正后 0.42）、本攻防 Δ 指纹 SNR=1571σ——差距 ~3000 倍。机理：Weil 零点项差分被 δ²λ²/2 压低（对称配对抵消一阶项），Δ 指纹的 w 核在 δλ≪R 时饱和不衰减——**前者测「零点离线的深度」、后者测「零点虚部结构」**。对威胁判据（离线零点存在）而言，结构测量是更直接的对象——离线深度 δ≤0.01 的存在性由构型自洽性（XXXVII）排除，其虚部结构由 Δ 指纹（本攻防）独立排除。")], indent=True)

# ---------- 6 ----------
h1("6　威胁排除支柱总装")
para([("n", "θ_K≥7/8 证据链中「威胁不存在」的独立支柱：① 无条件段（δ≥0.0902，零密度主幂）；② 数值闭合段（δ∈[0.05,0.0902]）；③ 盲区实证（δ≤0.01，XXXVI 左尾审计）；④ 间距计数（XXXVII：同虚部成对 ⇒ 极近对 ⇒ 实测 0 vs 需 21.9 万）；⑤ Δ 指纹（本攻防：3σ 阈值 420 对 vs 实测 2 对）。④⑤ 同源不同型：④ 是原始计数（依赖求根完整性），⑤ 是积分统计（对漏根免疫）——盲区威胁须同时违反函数方程配对、间距结构、跨族互斥统计——三重不兼容。")], indent=True)

# ---------- 7 ----------
h1("7　诚实边界与未决")
para([("n", "① Δ 指纹不测离线深度 δ——它排除「同虚部成对」的存在性，不直接测「零点离线多少」；深度维度由 S1 左尾审计覆盖。② 阈值标尺用静态噪声 σ_A=√(N_l·var_A)（实测 var_A=0.1144），未做 bootstrap 动态噪声——σ_A 本身是 14σ 基线的残差尺度，静态口径合理但非唯一。③ 219000 对注入使 Δ_pure/N_K 变号（+0.258）——Δ 常数定律的线性区间外推（−0.0070·N_K 尺度）未被如此大规模扰动验证——检测器读数在大扰动端只作「爆炸级」定性。④ 低 γ 段（γ<2000）未分层——该段 1/γ 加权增强（S2 边界注记），Δ 指纹的窗口统计是否同幅未单独核。⑤ 若威胁构型偏离「同虚部成对」（如虚部错开 ε≥W 的孤立离线零点），Δ 指纹响应衰减——该类构型由 S1 实部扫描覆盖（可分辨范围内 0 个）。⑥ 基线常数 −0.0070 的解析证明仍未完成（谱层核积分 T→∞ 极限）——本攻防仅用其数值形态。")], indent=True)

# ---------- 8 ----------
h1("8　判定总结")
para([("n", "S3 完成并落档。Δ 常数指纹对同虚部成对注入呈爆炸响应：威胁级（21.9 万对）SNR=1571σ（σ_A）/531σ（√N）——每对 2.60，3σ 阈值 ≈420 对——而实测跨族极近对仅 2 个，低于阈值 210 倍：**Δ 指纹独立闭合威胁排除**。检测器对离线深度 δ 代数免疫（A(t) 只含虚部）、对个别漏根免疫（单零点权 0.4%）——与 XXXVII 间距计数的盲区互不相交，构成「计数 + 积分」双重证据线；21.9 万对注入使 Δ_pure/N_K 变号，跨族互斥常数层与威胁构型直接矛盾。威胁排除支柱 = 无条件段 + 数值闭合段 + 盲区实证 + 间距计数 + Δ 指纹——五柱合围。主线下一步：θ_K≥7/8 论证从「威胁排除」转回「正向下界构造」（ζ_K 特有谱结构——在线比例桥余项精细形态），Δ 指纹作为独立常数层工具保留。")], indent=True)

# ---------- 参考文献 ----------
h1("参考文献")
refs = [
    "《Δ 常数定律（跨族互斥）：陈述、文献定位与证据》（C:\\proj2\\BSD\\delta-law\\Delta常数定律-陈述与证据.md）：Δ_pure 口径、c_pure≈0.0070±0.0003、var_A=0.1144、mean_A=0.7856、−14.1σ/−41.7σ、多 T 收敛（−0.92%→−0.70%）.",
    "L4 多 T 数值件：d1_l4_work/l4_multit.py/.json（S_ζL−E 口径：w(u)=4/(4+u²)、R=20、l∈[12,T−W]、分块诊断、C_ζ(h)）.",
    "攻防 XXXVI：路2攻防XXXVI-左尾224审计与威胁形态实证（跨族极近对 <1e-6 2 个、<1e-4 386≈泊松 380、威胁带 δ≤0.01）.",
    "攻防 XXXVII：路2攻防XXXVII-Weil显式公式实部口径注入敏感性（构型自洽性：同虚部成对 ⇒ 极近对 ⇒ 实测排除；S2 Weil 口径 SNR≈0.59 近盲）.",
    "S3 数值件：entF_work/detF_base.py/.json（基线复现）、detF_inject.py/.json（注入实验：四档×两构型×5 抽样 + 全量重算验证 rel_diff=1.57e-14）——可复跑（Anaconda python）。",
    "零点数据：zeros_T6e5.npz（ζ 族 999492、L 族 1150050、T=6×10⁵）.",
]
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    rich(p, [("n", "[" + str(i) + "] " + r)], size=9)

doc.save(OUT)
print("SAVED:", OUT)

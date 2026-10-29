# -*- coding: utf-8 -*-
"""arXiv 草稿改进：补 C_MT < 4/3 已 Lean 真证 + 更新 §9 计数（编辑场景，原地写回）。"""
import copy
import docx
from docx.oxml.ns import qn

TARGET = r'C:\Users\shtcl\Doubao\chats\2026-10-16\document\d1\en\QRH-arXiv-draft-single-family-pairing-constant-L(s,chi5).docx'
doc = docx.Document(TARGET)


def iter_wt(p):
    return [t for t in p._p.iter(qn('w:t'))]


def replace_in_wt(wts, old, new):
    n = 0
    for t in wts:
        if old in t.text:
            t.text = t.text.replace(old, new)
            n += 1
    return n


def add_run_after_last(p, text):
    """段尾追加 run，继承段落最后一个 w:r 的 rPr（若有）。"""
    r = p.add_run(text)
    # 继承字体：复制同段内最后一个 w:r 的 rPr
    last_rPr = None
    for el in p._p.findall(qn('w:r')):
        rPr = el.find(qn('w:rPr'))
        if rPr is not None:
            last_rPr = rPr
    if last_rPr is not None:
        new_rPr = copy.deepcopy(last_rPr)
        existing = r._r.find(qn('w:rPr'))
        if existing is not None:
            r._r.remove(existing)
        r._r.insert(0, new_rPr)
    return r


# ---- 段 13（摘要）：更新计数并提 C_MT ----
p13 = doc.paragraphs[13]
wt13 = iter_wt(p13)
c1 = replace_in_wt(wt13, '26 lemmas proved and a full build of 2604 jobs passing',
                   '35+ lemmas proved and a full build of 2704 jobs passing (2026-10-29), including the standing constant bound C_MT < 4/3')

# ---- 段 24（§3 变分机制）：C_MT 数值处补已证指向 ----
p24 = doc.paragraphs[24]
add_run_after_last(p24,
    ' The bound C_MT < 4/3 is proved in Lean (Section 9; pure rational chain, 2026-10-29).')

# ---- 段 58（§9 Lean 章节）：更新数字 + 段尾补 C_MT 真证细节 ----
p58 = doc.paragraphs[58]
wt58 = iter_wt(p58)
c2 = replace_in_wt(wt58, '2604 jobs', '2704 jobs')
c3 = replace_in_wt(wt58, '26 lemmas are proved', '35+ lemmas are proved')
c4 = replace_in_wt(wt58, '(2026-10-27, main_l2_full.txt)', '(2026-10-29, dbg_build_fourier.txt)')
add_run_after_last(p58,
    ' The standing constant bound C_MT < 4/3 is proved in Lean as C_MT_lt_four_thirds: the cos fourth-order upper bound 73/96 and the sin lower bound 11/(12*sqrt 2) - sqrt 2/800 at x0 = 1/sqrt 2 give cot(1/sqrt 2) < 1825/(1097*sqrt 2), hence C_MT < 2922/2194 < 4/3 - a pure rational chain with no integral layer and no aggregate dependency. The Fourier operator fourierCos is now a Bochner-integral true definition and its scale law fourierCos_scale is proved (2026-10-29); the build was rerun at 2704 jobs, all green (dbg_build_fourier.txt).')

doc.save(TARGET)
print('replacements: para13=%d para58_jobs=%d para58_lemmas=%d para58_build=%d' % (c1, c2, c3, c4))
print('saved:', TARGET)

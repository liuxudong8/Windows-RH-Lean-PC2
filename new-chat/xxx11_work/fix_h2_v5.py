# -*- coding: utf-8 -*-
"""攻防 XLI 修订 v5：§5 第二层末尾插入"外推说明"段（防有限 T 眼中钉）。
位置：中文"…未发现偏离拟随机框架的结构性证据。"段后；英文"…no structural evidence against the pseudorandom framework."段后。
格式：复制前段段落属性。"""
import copy, shutil
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ZH = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证.docx"
EN = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\en\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证-EN.docx"
WORK = r"C:\Users\shtcl\Doubao\chats\2026-10-16\new-chat\xxx11_work"

shutil.copy(ZH, WORK + r"\SOURCE_XLI_ZH.v4.docx")
shutil.copy(EN, WORK + r"\SOURCE_XLI_EN.v4.docx")
print("v4 snapshots saved")

def insert_para_after(para, text):
    new_el = copy.deepcopy(para._p)
    for r in list(new_el.findall(qn('w:r'))):
        new_el.remove(r)
    for h in list(new_el.findall(qn('w:hyperlink'))):
        new_el.remove(h)
    para._p.addnext(new_el)
    np = Paragraph(new_el, para._parent)
    np.add_run(text)
    return np

def find_para(doc, anchor):
    for p in doc.paragraphs:
        if anchor in p.text:
            return p
    return None

# ---------------- 中文版 ----------------
doc = Document(ZH)
zh_anchor = "未发现偏离拟随机框架的结构性证据"
p = find_para(doc, zh_anchor)
assert p is not None, "zh anchor not found"
zh_new = ("外推说明：我们在 T=6×10⁵ 的 5 个子区间内做了漂移检验，最优核 A 的相邻段间漂移 ≤0.005、"
          "全程单调无发散。前提（A≤1.125）的数值支撑因此是稳固的，其合法性仅依赖一个显式外推假设——"
          "无限高度上的全局均值不发生剧烈反转（该假设不可直接验证，可由更大 T 的数值检验强化）。")
insert_para_after(p, zh_new)
doc.save(ZH)
print("[zh] extrapolation paragraph inserted")

# ---------------- 英文版 ----------------
doc = Document(EN)
en_anchor = "finds no structural evidence against the pseudorandom framework"
p = find_para(doc, en_anchor)
assert p is not None, "en anchor not found"
en_new = ("Extrapolation statement: the drift check was performed on the 5 subintervals within T=6×10⁵, "
          "with adjacent-segment drift of the optimal-kernel A at most 0.005 and a monotone, non-divergent "
          "trend overall. The numerical support for the premise (A≤1.125) is accordingly stable; its "
          "legitimacy rests on one explicit extrapolation assumption — that the global mean over infinite "
          "height does not undergo a violent reversal (not directly verifiable; can be strengthened by "
          "numerical checks at larger T).")
insert_para_after(p, en_new)
doc.save(EN)
print("[en] extrapolation paragraph inserted")

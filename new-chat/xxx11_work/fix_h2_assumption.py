# -*- coding: utf-8 -*-
"""攻防 XLI 修订：将 L 族 BGS 型配对公式存在性提升为显式假设 (H2)，中英两版同步。"""
import re
from docx import Document

FILES = {
"zh": r"D:\F\有名公司同事文档\刘旭东文集\document\d1\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证.docx",
"en": r"D:\F\有名公司同事文档\刘旭东文集\document\d1\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证-EN.docx",
}

REPL = {
"zh": {
# 摘要尾部
"L 族配对公式存在性在文献中无无条件版（BGS 推广为框架内开放构造）——定理形态定为纯假设（弱前提）。":
"L 族配对公式存在性在文献中无无条件版（BGS 推广为框架内开放构造）——定理以显式假设 (H2) 覆盖 L 族 BGS 型公式的存在性，与 A_L≤1.125 并列，定理形态定为纯假设（弱前提）。",
# §5 第一层
"条件定理前提从「GUE 配对相关（整体、不可验证）」降为「A_ζ,A_L≤1.125（配对常数假设）」——GUE 是充分非必要，配对常数是 audit6 直接数值验证的对象。A≤1.125 门槛数值全段通过（最优核 1.008-1.014）。":
"条件定理前提从「GUE 配对相关（整体、不可验证）」降为显式假设 (H1)-(H2)：A_ζ≤1.125，且 L 族满足 BGS 型显式公式并 A_L≤1.125——GUE 是充分非必要，配对常数是 audit6 直接数值验证的对象；(H2) 的公式存在性独立于数值，属文献开放构造（见 §7②）。A≤1.125 门槛数值全段通过（最优核 1.008-1.014）。",
# §6 定理
"定理（弱前提拟随机并集定理，条件形态）：设 ζ 与 L(χ₅) 的零点在归一化尺度上的配对常数满足 A_ζ,A_L≤1.125（拟随机前提，含 L 族配对公式存在性假设；数值最优核实测 A_ζ=1.0079、A_L=1.0114）。则并集在线简单零点比例 ≥ (N_ζ(2−A_ζ)+N_L(2−A_L))/N_K ≥ 7/8。":
"定理（弱前提拟随机并集定理，条件形态）：设如下两个显式假设成立——(H1) ζ 的零点在归一化尺度上的配对常数 A_ζ≤1.125；(H2) L(s,χ₅) 的单族配对相关满足 BGS 型显式公式（即 BGS 无条件配对定理对固定模实特征 χ₅ 的推广成立；这是 L 族配对和 ΣK²_L 有定义且可受控的前提），且其配对常数 A_L≤1.125（数值最优核实测 A_ζ=1.0079、A_L=1.0114）。则并集在线简单零点比例 ≥ (N_ζ(2−A_ζ)+N_L(2−A_L))/N_K ≥ 7/8。",
# §6 证明骨架 (iii)
"证明骨架：(i) N₀ˢ_K=N₀ˢ_ζ+N₀ˢ_L（零点互异，已核实）；(ii) 两族分别应用 Proposition 2.1：N₀ˢ_ζ≥2N_ζ−ΣK²_ζ、N₀ˢ_L≥2N_L−ΣK²_L（无条件不等式，K=η²^、K(0)=1）；(iii) 拟随机前提 ⟹ ΣK²_ζ≤A_ζN_ζ、ΣK²_L≤A_LN_L；(iv) 代入加权即得。分解论证的关键：跨族配对完全不进入（与并集直接论证的 A_union=1.29-1.62 判负形成对照）。":
"证明骨架：(i) N₀ˢ_K=N₀ˢ_ζ+N₀ˢ_L（零点互异，已核实）；(ii) 两族分别应用 Proposition 2.1：N₀ˢ_ζ≥2N_ζ−ΣK²_ζ、N₀ˢ_L≥2N_L−ΣK²_L（无条件不等式，K=η²^、K(0)=1）；(iii) (H1)-(H2) ⟹ ΣK²_ζ≤A_ζN_ζ、ΣK²_L≤A_LN_L（ΣK²_L 在 (H2) 的 BGS 型公式下定义）；(iv) 代入加权即得。分解论证的关键：跨族配对完全不进入（与并集直接论证的 A_union=1.29-1.62 判负形成对照）。",
# §7 ②
"② L 族配对公式存在性在文献中无无条件版（Özlük 需 GRH、Yıldırım 需 GRH+跨字符平均）——BGS 推广至固定模实字符为框架内开放构造，定理以假设覆盖。":
"② L 族配对公式存在性在文献中无无条件版（Özlük 需 GRH、Yıldırım 需 GRH+跨字符平均）——BGS 推广至固定模实字符为框架内开放构造，故定理将其列为显式假设 (H2)（L 族满足 BGS 型显式公式），与数值前提 A_L≤1.125 分离：没有 (H2)，A_L 甚至无定义。",
},
"en": {
# Abstract 尾部
"No unconditional version of the pair-correlation formula for the L family exists in the literature (the BGS generalization is an open construction within the framework) — the theorem takes the purely conditional (weak-premise) form.":
"No unconditional version of the pair-correlation formula for the L family exists in the literature (the BGS generalization is an open construction within the framework) — the theorem covers the existence of the BGS-type formula for the L family by the explicit hypothesis (H2), stated alongside A_L≤1.125, and takes the purely conditional (weak-premise) form.",
# Tier 1
"The conditional-theorem premise is downgraded from \u201cGUE pair correlation (global, unverifiable)\u201d to \u201cA_ζ, A_L≤1.125 (pair-correlation-constant hypothesis)\u201d — GUE is sufficient but not necessary; the pair constant is the object directly verified numerically by audit6. The A≤1.125 threshold passes numerically over the whole range (optimal kernel 1.008–1.014).":
"The conditional-theorem premise is downgraded from \u201cGUE pair correlation (global, unverifiable)\u201d to the explicit hypotheses (H1)–(H2): A_ζ≤1.125, and the L family satisfying a BGS-type explicit formula with A_L≤1.125 — GUE is sufficient but not necessary; the pair constants are the objects directly verified numerically by audit6, while the formula existence in (H2) is independent of numerics and is an open construction in the literature (see §7 ②). The A≤1.125 threshold passes numerically over the whole range (optimal kernel 1.008–1.014).",
# Theorem
"Theorem (weak-premise pseudorandom union theorem, conditional form). Let the pair-correlation constants of the zeros of ζ and L(χ₅) on the normalized scale satisfy A_ζ, A_L≤1.125 (pseudorandom premise, including the existence hypothesis of the pair-correlation formula for the L family; measured optimal kernels A_ζ=1.0079, A_L=1.0114). Then the union online simple-zero proportion is ≥ (N_ζ(2−A_ζ)+N_L(2−A_L))/N_K ≥ 7/8.":
"Theorem (weak-premise pseudorandom union theorem, conditional form). Assume the following two explicit hypotheses: (H1) the pair-correlation constant of the zeros of ζ on the normalized scale satisfies A_ζ≤1.125; (H2) the single-family pair correlation of L(s,χ₅) satisfies a BGS-type explicit formula (i.e., the unconditional BGS pair-correlation theorem generalized to the fixed real character χ₅ — the hypothesis under which the L-family pair sum ΣK²_L is defined and controllable), and its pair-correlation constant satisfies A_L≤1.125 (measured optimal kernels A_ζ=1.0079, A_L=1.0114). Then the union online simple-zero proportion is ≥ (N_ζ(2−A_ζ)+N_L(2−A_L))/N_K ≥ 7/8.",
# Proof skeleton (iii)
"Proof skeleton: (i) N₀ˢ_K=N₀ˢ_ζ+N₀ˢ_L (zeros distinct, verified); (ii) apply Proposition 2.1 separately: N₀ˢ_ζ≥2N_ζ−ΣK²_ζ, N₀ˢ_L≥2N_L−ΣK²_L (unconditional inequalities, K=η²^, K(0)=1); (iii) the pseudorandom premise ⟹ ΣK²_ζ≤A_ζN_ζ, ΣK²_L≤A_LN_L; (iv) substitute and weight. The key point of the decomposition argument: cross-family pairings never enter (in contrast with the failing direct union argument, A_union=1.29–1.62).":
"Proof skeleton: (i) N₀ˢ_K=N₀ˢ_ζ+N₀ˢ_L (zeros distinct, verified); (ii) apply Proposition 2.1 separately: N₀ˢ_ζ≥2N_ζ−ΣK²_ζ, N₀ˢ_L≥2N_L−ΣK²_L (unconditional inequalities, K=η²^, K(0)=1); (iii) (H1)–(H2) ⟹ ΣK²_ζ≤A_ζN_ζ, ΣK²_L≤A_LN_L (ΣK²_L being defined under the BGS-type formula of (H2)); (iv) substitute and weight. The key point of the decomposition argument: cross-family pairings never enter (in contrast with the failing direct union argument, A_union=1.29–1.62).",
# §7 ②
"\u2461 No unconditional version of the L-family pair-correlation formula exists in the literature (Özlük needs GRH; Yıldırım needs GRH + cross-character average) — the BGS generalization to a fixed-modulus real character is an open construction within the framework, covered by hypothesis in the theorem.":
"\u2461 No unconditional version of the L-family pair-correlation formula exists in the literature (Özlük needs GRH; Yıldırım needs GRH + cross-character average) — the BGS generalization to a fixed-modulus real character is an open construction within the framework, so the theorem states it as the explicit hypothesis (H2) (the L family satisfies a BGS-type formula), separated from the numerical premise A_L≤1.125: without (H2), A_L is not even defined.",
},
}

def norm(s):
    return re.sub(r"\s+", " ", s.strip())

def set_para_text(p, new_text):
    runs = p.runs
    if not runs:
        p.add_run(new_text)
        return
    first = runs[0]
    first.text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

for lang, path in FILES.items():
    doc = Document(path)
    repl = {norm(k): v for k, v in REPL[lang].items()}
    hit, miss = 0, []
    for p in doc.paragraphs:
        key = norm(p.text)
        if key in repl:
            set_para_text(p, repl[key])
            hit += 1
        else:
            # 只对包含关键片的段落做诊断（防错）
            for frag in ("(H2)", "pair-correlation formula for the L family", "BGS 推广为框架内开放构造"):
                if frag in p.text and key not in repl:
                    miss.append(key[:60])
                    break
    doc.save(path)
    print(f"[{lang}] replaced: {hit}, suspicious unmatched: {len(miss)}")
    for m in miss[:10]:
        print("   MISS?", m)

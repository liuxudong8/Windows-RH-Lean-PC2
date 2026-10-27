# -*- coding: utf-8 -*-
"""攻防 XLI 修订 v3：定理钉核 + 骨架(iii)恒等式归定义 + §7②"无定义"措辞修正。中英同步。"""
import shutil
from docx import Document

ZH = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证.docx"
EN = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\en\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证-EN.docx"
WORK = r"C:\Users\shtcl\Doubao\chats\2026-10-16\new-chat\xxx11_work"

shutil.copy(ZH, WORK + r"\SOURCE_XLI_ZH.v2.docx")
shutil.copy(EN, WORK + r"\SOURCE_XLI_EN.v2.docx")
print("v2 snapshots saved")

def set_para_text(p, new_text):
    runs = p.runs
    if not runs:
        p.add_run(new_text); return
    runs[0].text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

def para_replace(doc, anchor, repl, whole=False):
    for p in doc.paragraphs:
        if anchor in p.text:
            if whole:
                set_para_text(p, repl)
            else:
                set_para_text(p, p.text.replace(anchor, repl))
            return True
    return False

# ================= 中文版 =================
doc = Document(ZH)

# 定理整段：钉核（存在性版本） + (H2) 括注改"锚定到 BGS 型理论坐标"
zh_thm_old = "定理（弱前提拟随机并集定理，条件形态）：设如下两个显式假设成立——(H1) ζ 的零点在归一化尺度上的配对常数 A_ζ≤1.125；(H2) L(s,χ₅) 的单族配对相关满足 BGS 型显式公式（即 BGS 无条件配对定理对固定模实特征 χ₅ 的推广成立；这是 L 族配对和 ΣK²_L 有定义且可受控的前提），且其配对常数 A_L≤1.125（数值最优核实测 A_ζ=1.0079、A_L=1.0114）。则并集在线简单零点比例 ≥ (N_ζ(2−A_ζ)+N_L(2−A_L))/N_K ≥ 7/8。"
zh_thm_new = "定理（弱前提拟随机并集定理，条件形态）：设如下两个显式假设成立——(H1) 存在核 K₁（audit6 最优核族内）使 ζ 的零点在归一化尺度上的配对常数 A_ζ(K₁)≤1.125；(H2) L(s,χ₅) 的单族配对相关满足 BGS 型显式公式（即 BGS 无条件配对定理对固定模实特征 χ₅ 的推广成立；这是把 L 族配对和 ΣK²_L 锚定到 BGS 型理论坐标、使其可受控的前提），且存在核 K₂ 使 A_L(K₂)≤1.125（数值口径：audit6 最优核 fejer_a0.5，A_ζ=1.0079、A_L=1.0114）。则并集在线简单零点比例 ≥ (N_ζ(2−A_ζ(K₁))+N_L(2−A_L(K₂)))/N_K ≥ 7/8。"

# 骨架整段：恒等式归定义，(H1)(H2) 只负责门槛代入
zh_sk_old = "证明骨架：(i) N₀ˢ_K=N₀ˢ_ζ+N₀ˢ_L（零点互异，已核实）；(ii) 两族分别应用 Proposition 2.1：N₀ˢ_ζ≥2N_ζ−ΣK²_ζ、N₀ˢ_L≥2N_L−ΣK²_L（无条件不等式，K=η²^、K(0)=1）；(iii) (H1)-(H2) ⟹ ΣK²_ζ≤A_ζN_ζ、ΣK²_L≤A_LN_L（ΣK²_L 在 (H2) 的 BGS 型公式下定义）；(iv) 代入加权即得。分解论证的关键：跨族配对完全不进入（与并集直接论证的 A_union=1.29-1.62 判负形成对照）。"
zh_sk_new = "证明骨架：(i) N₀ˢ_K=N₀ˢ_ζ+N₀ˢ_L（零点互异，已核实）；(ii) 两族分别应用 Proposition 2.1：N₀ˢ_ζ≥2N_ζ−ΣK²_ζ、N₀ˢ_L≥2N_L−ΣK²_L（无条件不等式，K=η²^、K(0)=1）；(iii) 按定义 A_ζ=ΣK²_ζ/N_ζ、A_L=ΣK²_L/N_L（A_L 的受控坐标由 (H2) 的 BGS 型公式提供），则 Prop 2.1 给出 N₀ˢ_ζ≥(2−A_ζ)N_ζ、N₀ˢ_L≥(2−A_L)N_L；(iv) 代入前提 (H1)(H2)（A_ζ,A_L≤1.125）加权即得。分解论证的关键：跨族配对完全不进入（与并集直接论证的 A_union=1.29-1.62 判负形成对照）。"

# §7② 片段："无定义"措辞修正
zh_7_old = "故定理将其列为显式假设 (H2)（L 族满足 BGS 型显式公式），与数值前提 A_L≤1.125 分离：没有 (H2)，A_L 甚至无定义。"
zh_7_new = "故定理将其列为显式假设 (H2)（L 族满足 BGS 型显式公式），与数值前提 A_L≤1.125 分离：没有 (H2)，A_L≤1.125 就只是数值数据、没有 BGS 型理论坐标（受控配对常数无定义）。"

ok = {}
ok["zh_thm"] = para_replace(doc, zh_thm_old, zh_thm_new, whole=True)
ok["zh_sk"]  = para_replace(doc, zh_sk_old, zh_sk_new, whole=True)
ok["zh_s7"]  = para_replace(doc, zh_7_old, zh_7_new)
doc.save(ZH)
print("[zh]", ok)

# ================= 英文版 =================
doc = Document(EN)

en_thm_old = "Theorem (weak-premise pseudorandom union theorem, conditional form). Assume the following two explicit hypotheses: (H1) the pair-correlation constant of the zeros of ζ on the normalized scale satisfies A_ζ≤1.125; (H2) the single-family pair correlation of L(s,χ₅) satisfies a BGS-type explicit formula (i.e., the unconditional BGS pair-correlation theorem generalized to the fixed real character χ₅ — the hypothesis under which the L-family pair sum ΣK²_L is defined and controllable), and its pair-correlation constant satisfies A_L≤1.125 (measured optimal kernels A_ζ=1.0079, A_L=1.0114). Then the union online simple-zero proportion is ≥ (N_ζ(2−A_ζ)+N_L(2−A_L))/N_K ≥ 7/8."
en_thm_new = "Theorem (weak-premise pseudorandom union theorem, conditional form). Assume the following two explicit hypotheses: (H1) there exists a kernel K₁ (within the audit6 optimal-kernel family) such that the pair-correlation constant of the zeros of ζ on the normalized scale satisfies A_ζ(K₁)≤1.125; (H2) the single-family pair correlation of L(s,χ₅) satisfies a BGS-type explicit formula (i.e., the unconditional BGS pair-correlation theorem generalized to the fixed real character χ₅ — the BGS-type theoretical coordinate in which the L-family pair constant A_L is anchored and controllable), and there exists a kernel K₂ such that A_L(K₂)≤1.125 (numerical convention: audit6 optimal kernel fejer_a0.5, A_ζ=1.0079, A_L=1.0114). Then the union online simple-zero proportion is ≥ (N_ζ(2−A_ζ(K₁))+N_L(2−A_L(K₂)))/N_K ≥ 7/8."

en_sk_old = "Proof skeleton: (i) N₀ˢ_K=N₀ˢ_ζ+N₀ˢ_L (zeros distinct, verified); (ii) apply Proposition 2.1 separately: N₀ˢ_ζ≥2N_ζ−ΣK²_ζ, N₀ˢ_L≥2N_L−ΣK²_L (unconditional inequalities, K=η²^, K(0)=1); (iii) (H1)–(H2) ⟹ ΣK²_ζ≤A_ζN_ζ, ΣK²_L≤A_LN_L (ΣK²_L being defined under the BGS-type formula of (H2)); (iv) substitute and weight. The key point of the decomposition argument: cross-family pairings never enter (in contrast with the failing direct union argument, A_union=1.29–1.62)."
en_sk_new = "Proof skeleton: (i) N₀ˢ_K=N₀ˢ_ζ+N₀ˢ_L (zeros distinct, verified); (ii) apply Proposition 2.1 separately: N₀ˢ_ζ≥2N_ζ−ΣK²_ζ, N₀ˢ_L≥2N_L−ΣK²_L (unconditional inequalities, K=η²^, K(0)=1); (iii) by definition A_ζ=ΣK²_ζ/N_ζ and A_L=ΣK²_L/N_L (A_L controlled in the BGS-type coordinate of (H2)); Proposition 2.1 then gives N₀ˢ_ζ≥(2−A_ζ)N_ζ and N₀ˢ_L≥(2−A_L)N_L; (iv) substitute the premises (H1)–(H2) (A_ζ, A_L≤1.125) and weight. The key point of the decomposition argument: cross-family pairings never enter (in contrast with the failing direct union argument, A_union=1.29–1.62)."

en_7_old = "so the theorem states it as the explicit hypothesis (H2) (the L family satisfies a BGS-type formula), separated from the numerical premise A_L≤1.125: without (H2), A_L is not even defined."
en_7_new = "so the theorem states it as the explicit hypothesis (H2) (the L family satisfies a BGS-type formula), separated from the numerical premise A_L≤1.125: without (H2), A_L≤1.125 is merely numerical data, with no BGS-type theoretical coordinate (the controlled pair constant is undefined)."

ok = {}
ok["en_thm"] = para_replace(doc, en_thm_old, en_thm_new, whole=True)
ok["en_sk"]  = para_replace(doc, en_sk_old, en_sk_new, whole=True)
ok["en_s7"]  = para_replace(doc, en_7_old, en_7_new)
doc.save(EN)
print("[en]", ok)

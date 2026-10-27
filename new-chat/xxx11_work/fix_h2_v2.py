# -*- coding: utf-8 -*-
"""攻防 XLI 修订 v2：显式假设 (H2) 全文贯通。
中文版（用户已写 §5/§6/骨架）：补摘要尾部 + §7② 呼应；
英文版（d1\\en）：完整对齐 5 处。"""
import re, shutil
from docx import Document

ZH = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证.docx"
EN = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\en\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证-EN.docx"
WORK = r"C:\Users\shtcl\Doubao\chats\2026-10-16\new-chat\xxx11_work"

# 修改前快照（只读 baseline，供 audit --source 与回滚）
shutil.copy(ZH, WORK + r"\SOURCE_XLI_ZH.bak.docx")
shutil.copy(EN, WORK + r"\SOURCE_XLI_EN.bak.docx")
print("baselines saved")

# (段锚子串, 新整段文本) 或 (片段, 新片段)
def set_para_text(p, new_text):
    runs = p.runs
    if not runs:
        p.add_run(new_text); return
    first = runs[0]
    first.text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

def para_replace(doc, anchor, repl, whole=False):
    for p in doc.paragraphs:
        if anchor in p.text:
            if whole:
                set_para_text(p, repl)
            else:
                set_para_text(p, p.text.replace(anchor, repl))
            return True
    return False

# ---------------- 中文版：补摘要尾 + §7② ----------------
doc = Document(ZH)
zh_abs_old = "（BGS 推广为框架内开放构造）——定理形态定为纯假设（弱前提）。"
zh_abs_new = "（BGS 推广为框架内开放构造）——定理以显式假设 (H2) 覆盖 L 族 BGS 型公式的存在性，与 A_L≤1.125 并列，定理形态定为纯假设（弱前提）。"
zh_7_old = "BGS 推广至固定模实字符为框架内开放构造，定理以假设覆盖。"
zh_7_new = "BGS 推广至固定模实字符为框架内开放构造，故定理将其列为显式假设 (H2)（L 族满足 BGS 型显式公式），与数值前提 A_L≤1.125 分离：没有 (H2)，A_L 甚至无定义。"
ok1 = para_replace(doc, zh_abs_old, zh_abs_new)
ok2 = para_replace(doc, zh_7_old, zh_7_new)
doc.save(ZH)
print("[zh] abstract tail:", ok1, "| §7②:", ok2)

# ---------------- 英文版：5 处完整对齐 ----------------
doc = Document(EN)
# 1) 摘要尾部
en_abs_old = "(the BGS generalization is an open construction within the framework) — the theorem takes the purely conditional (weak-premise) form."
en_abs_new = "(the BGS generalization is an open construction within the framework) — the theorem covers the existence of the BGS-type formula for the L family by the explicit hypothesis (H2), stated alongside A_L≤1.125, and takes the purely conditional (weak-premise) form."
# 2) Tier 1 整段
en_t1_new = ("The conditional-theorem premise is downgraded from \u201cGUE pair correlation (global, unverifiable)\u201d "
             "to the explicit hypotheses (H1)–(H2): A_ζ≤1.125, and the L family satisfying a BGS-type explicit formula "
             "with A_L≤1.125 — GUE is sufficient but not necessary; the pair constants are the objects directly verified "
             "numerically by audit6, while the formula existence in (H2) is independent of numerics and is an open "
             "construction in the literature (see §7 ②). The A≤1.125 threshold passes numerically over the whole range "
             "(optimal kernel 1.008–1.014).")
# 3) 定理整段
en_th_new = ("Theorem (weak-premise pseudorandom union theorem, conditional form). Assume the following two explicit "
             "hypotheses: (H1) the pair-correlation constant of the zeros of ζ on the normalized scale satisfies A_ζ≤1.125; "
             "(H2) the single-family pair correlation of L(s,χ₅) satisfies a BGS-type explicit formula (i.e., the "
             "unconditional BGS pair-correlation theorem generalized to the fixed real character χ₅ — the hypothesis under "
             "which the L-family pair sum ΣK²_L is defined and controllable), and its pair-correlation constant satisfies "
             "A_L≤1.125 (measured optimal kernels A_ζ=1.0079, A_L=1.0114). Then the union online simple-zero proportion "
             "is ≥ (N_ζ(2−A_ζ)+N_L(2−A_L))/N_K ≥ 7/8.")
# 4) 证明骨架 (iii)
en_sk_old = "(iii) the pseudorandom premise ⟹ ΣK²_ζ≤A_ζN_ζ, ΣK²_L≤A_LN_L;"
en_sk_new = "(iii) (H1)–(H2) ⟹ ΣK²_ζ≤A_ζN_ζ, ΣK²_L≤A_LN_L (ΣK²_L being defined under the BGS-type formula of (H2));"
# 5) §7 ②
en_7_old = "within the framework, covered by hypothesis in the theorem."
en_7_new = ("within the framework, so the theorem states it as the explicit hypothesis (H2) (the L family satisfies a "
            "BGS-type formula), separated from the numerical premise A_L≤1.125: without (H2), A_L is not even defined.")

ok = {}
ok["abs"]  = para_replace(doc, en_abs_old, en_abs_new)
ok["t1"]   = para_replace(doc, "The conditional-theorem premise is downgraded from", en_t1_new, whole=True)
ok["thm"]  = para_replace(doc, "Theorem (weak-premise pseudorandom union theorem, conditional form)", en_th_new, whole=True)
ok["skel"] = para_replace(doc, en_sk_old, en_sk_new)
ok["s7"]   = para_replace(doc, en_7_old, en_7_new)
doc.save(EN)
print("[en]", ok)

# -*- coding: utf-8 -*-
"""攻防 XLI 修订 v6：§8 主线候选句更新——"在线比例→1（强配对全尺度）"立为下一阶段主线；
不引入 31/32；7/8 保持现行定理坐标。中英同步，保留段落首 run 格式。"""
import copy, shutil
from docx import Document
from docx.oxml.ns import qn

ZH = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证.docx"
EN = r"D:\F\有名公司同事文档\刘旭东文集\document\d1\en\路2攻防XLI-拟随机并集定理的常数机制闭合-QRH攻坚0.6725到7-8的最优核数值构造与前提可靠性三层论证-EN.docx"
WORK = r"C:\Users\shtcl\Doubao\chats\2026-10-16\new-chat\xxx11_work"

shutil.copy(ZH, WORK + r"\SOURCE_XLI_ZH.v5.docx")
shutil.copy(EN, WORK + r"\SOURCE_XLI_EN.v5.docx")
print("v5 snapshots saved")

def replace_para_text(doc, anchor, new_text):
    for p in doc.paragraphs:
        if anchor in p.text:
            runs = p.runs
            rpr = None
            if runs:
                rpr_el = runs[0]._r.find(qn('w:rPr'))
                if rpr_el is not None:
                    rpr = copy.deepcopy(rpr_el)
            for r in list(p._p.findall(qn('w:r'))):
                p._p.remove(r)
            for h in list(p._p.findall(qn('w:hyperlink'))):
                p._p.remove(h)
            new_run = p.add_run(new_text)
            if rpr is not None:
                new_run._r.insert(0, rpr)
            return True
    return False

# ---------------- 中文版 ----------------
doc = Document(ZH)
zh_anchor = "主线下一步候选"
zh_new = ("主线下一步候选：(a) 分解论证严格化（两族 Proposition 2.1 的有限 T 版本）；"
          "(b) L 族配对公式的无条件推广构造（BGS 方法适配模 5 实字符）；"
          "(c) 下一阶段主线——转回 ζ_K 特有谱结构攻坚（在线比例桥余项精细形态——超密度假设级专用论证），"
          "终极形态为在线简单零点比例 1−o(1)：将 (H2) 升级为 λ(T)→∞ 的全尺度 GUE 配对"
          "（强 Montgomery 型对固定模实特征 L(s,χ₅)，框架内开放构造），"
          "则 Prop 2.1 直接给出配对常数 A→1、在线比例 →1——即文献「强配对猜想 ⟹ 几乎全部零点简单」的框架内版本；"
          "7/8 门槛保持为现行定理坐标。")
ok = replace_para_text(doc, zh_anchor, zh_new)
assert ok, "zh anchor not found"
doc.save(ZH)
print("[zh] main-line sentence updated")

# ---------------- 英文版 ----------------
doc = Document(EN)
en_anchor = "Next-step candidates"
en_new = ("Next-step candidates: (a) strictification of the decomposition argument (finite-T versions of "
          "Proposition 2.1 for the two families); (b) unconditional generalization of the L-family "
          "pair-correlation formula (BGS method adapted to modulus-5 real characters); (c) the next-stage "
          "main line — return to the ζ_K-specific spectral-structure push (fine form of the "
          "online-proportion-bridge remainder — a specialized argument at the super-density-hypothesis "
          "level), whose terminal form is an online simple-zero proportion of 1−o(1): upgrade (H2) to "
          "full-scale GUE pair correlation with λ(T)→∞ (a strong-Montgomery-type statement for L(s,χ₅) of "
          "fixed real character — an open construction within the framework); Proposition 2.1 then gives "
          "directly A→1 and online proportion →1 — the in-framework version of the literature’s "
          "“strong pair-correlation conjecture ⟹ almost all zeros simple”; the 7/8 threshold remains the "
          "current theorem coordinate.")
ok = replace_para_text(doc, en_anchor, en_new)
assert ok, "en anchor not found"
doc.save(EN)
print("[en] main-line sentence updated")

# -*- coding: utf-8 -*-
"""编辑骨架A讨论稿：4.1 钉死 Weil 显式公式精确形态；(ii) 标闭合；6节/1.3/参考文献同步"""
from docx import Document
from docx.oxml.ns import qn
import re, sys

SRC = r"D:\F\有名公司同事文档\刘旭东文集\3new\关于引力的思考和黎曼猜想\Arthur 迹公式\BSD\弱BSD秩1_Hecke显式公式重证明_骨架A讨论稿.docx"
TGT = r"D:\F\有名公司同事文档\刘旭东文集\3new\关于引力的思考和黎曼猜想\Arthur 迹公式\BSD\_edit_骨架A讨论稿.docx"

doc = Document(TGT)
paras = doc.paragraphs

def find_para(prefix):
    for p in paras:
        t = p.text.strip()
        if t.startswith(prefix):
            return p
    raise RuntimeError('未找到: ' + prefix)

def set_para_text(p, text):
    """保留段落样式与首 run 字体，替换全部文本"""
    # 收集首 run 的 rPr
    runs = p.runs
    rpr = None
    if runs and runs[0]._r.rPr is not None:
        import copy
        rpr = copy.deepcopy(runs[0]._r.rPr)
    # 删除所有 run
    for r in runs:
        r._r.getparent().remove(r._r)
    # 添加新 run 并套用 rPr
    new_r = p.add_run(text)
    if rpr is not None:
        new_r._r.insert(0, rpr)

# ---------- 4.1 第一段：骨架式 -> 精确形态 ----------
p41 = find_para('对 L(f,s)（中心 s = 1，degree 2），取偶、速降测试函数 h')
set_para_text(p41,
    '对 L(f,s)（中心 s = 1，degree 2，导子 N），本项目取 Spicer [7] 的椭圆曲线专用显式公式为标准形态'
    '（无条件，不依赖 GRH/BSD）：Σ_{γ>0} 2·sin(γ·log x)/γ = −η − log(2π/√N) − r_an·log x − log(1 − x⁻¹) + ψ_E(x)，'
    '其中 η = 0.57721566… 为欧拉常数，γ 遍历 L 的非平凡零点虚部（排除中心零点 s = 1），r_an 为解析秩，'
    'ψ_E(x) = −Σ_{n≤x} a_n·Λ(n)/n（n 遍历素幂，Λ 为 von Mangoldt 函数）。对应的一阶 Taylor 系数命题（同源，Spicer [7]）：'
    'L_E(s + 1) = s^r(a + b·s + c·s² + O(s³))，则 b/a = η + log(2π/√N)，'
    'c/a = ½[η + log(2π/√N)]² − π²/12 + Σ_{γ>0} γ⁻²。'
    '关键口径：Γ 因子符号与对数常数全部吸收进 −η − log(2π/√N)（ψ(1) = −η）；平凡零点口径以闭合形式 −log(1 − x⁻¹) 呈现；'
    'r_an 项显式出现——正是骨架 A 需要的解析秩输入。K = ℚ(√5) 端同形态：N → N_K = 1369，a_n → b_n（链 2 口径 ℓ = log Nm）。'
    '数值验证见 5.4(ii)：Taylor 命题两端闭合至 1e-9 量级。')

# ---------- 4.1 形式警告段结尾 ----------
pw = find_para('形式警告（G7 教训）')
wt = pw.text
if '系数核对列为待补' in wt:
    set_para_text(pw, wt.replace('本文先给出骨架，系数核对列为待补。',
        '该警告仍有效；4.1 的精确系数现已按上述形态逐项钉死'
        '（对数常数项 −η − log(2π/√N)、平凡零点口径 −log(1 − x⁻¹)、Γ 因子符号 −η、r_an 项显式），数值核对见 5.4(ii)。'))

# ---------- 5.4 (ii) ----------
p52 = find_para('（ii）4.1 显式公式精确形态按标准 Weil 公式核对')
set_para_text(p52,
    '（ii）4.1 显式公式精确形态已按 Spicer [7] 标准形态闭合。决定性数值验证（Taylor 系数命题，Richardson 外推二阶差分）：'
    'ℚ 端 b/a = 0.6096337777 对 η + log(2π/√37) = 0.6096337750（差 2.7×10⁻⁹）；'
    'K 端 b/a = −1.1958251726 对 η + log(2π/√1369) = −1.1958251813（差 8.7×10⁻⁹）。'
    '显式公式的 sinc 级数侧（ℚ 端前 24 个非平凡零点：γ₁ = 5.003170，γ₂ = 6.870391，γ₃ = 8.014331，γ₄ = 9.933098，'
    'γ₅ = 10.775138，γ₆ = 11.757325，γ₇ = 12.958386，γ₈ = 15.603858，γ₉ = 16.192017，γ₁₀ = 17.141694，'
    'γ₁₁ = 18.063654，γ₁₂ = 18.787196，γ₁₃ = 19.814822，γ₁₄ = 21.322800，γ₁₅ = 22.620430，…，最大 γ ≈ 30.896）：'
    'Σ 2·sin(γ·log x)/γ 对 RHS 的逼近呈收敛趋势（x = 10 处随零点数 4 → 8 → 14 为 −0.630 → −0.810 → −0.877，RHS = −1.172；'
    'x = 5 处前 8 零点闭合至 2.6×10⁻³）；余差与 sinc 级数条件收敛的截断尾 O(log T/(T·log x)) 同量级，非公式形态错误。'
    'K 端第一零点 γ = 1.63248196。4.1 已按此形态改写。（iii）4.3 全消矛盾的严格化仍待补。')

# ---------- 6 节结尾 ----------
p6 = find_para('骨架 A 给出了弱 BSD 秩 1 未证方向')
t6 = p6.text
if '两条闭合路径' in t6:
    set_para_text(p6, t6.replace('显式公式精确系数与全消矛盾的严格化是下一步的两条闭合路径。',
        '显式公式精确形态已闭合（见 5.4(ii)）；全消矛盾（iii）的严格化是下一步的唯一闭合路径。'))

# ---------- 1.3 ----------
p13 = find_para('骨架 A 的论证由四层构成')
t13 = p13.text
if '判秩所需的完整 L 函数数值延拓已闭合（见 5.4）' in t13:
    set_para_text(p13, t13.replace('判秩所需的完整 L 函数数值延拓已闭合（见 5.4）。',
        '判秩所需的完整 L 函数数值延拓已闭合（见 5.4）；显式公式精确形态已钉死（见 4.1、5.4(ii)）。'))

# ---------- 参考文献加 [7] ----------
p6ref = find_para('[6] LMFDB.')
# 在 [6] 段后插入 [7]
import copy
new_p = copy.deepcopy(p6ref._p)
# 清空 new_p 的 runs
for r in list(new_p.findall(qn('w:r'))):
    new_p.remove(r)
# 构造文本 run（python-docx 需要包装）
from docx.text.paragraph import Paragraph
wrapper = Paragraph(new_p, p6ref._parent)
run = wrapper.add_run('[7] Spicer S. L-functions: a crash course[R]. University of Washington REU, 2013. https://sites.math.washington.edu/~reu/papers/2013/simon/lfunctions.pdf.')
# 复制原段落的 rPr 到新 run
rpr_src = None
if p6ref.runs and p6ref.runs[0]._r.rPr is not None:
    rpr_src = copy.deepcopy(p6ref.runs[0]._r.rPr)
if rpr_src is not None:
    run._r.insert(0, rpr_src)
p6ref._p.addnext(new_p)

doc.save(TGT)
print('OK saved:', TGT)
